import os
import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load the master updated Excel
source_file = 'All_Website_Pages_SEO_Audit_UPDATED.xlsx'
wb_source = openpyxl.load_workbook(source_file, data_only=False)
ws_source = wb_source['All Website Pages SEO Audit']

total_data_rows = ws_source.max_row - 1
total_cols = ws_source.max_column
print(f"Total pages in master audit: {total_data_rows}")

# Chunk size
chunk_size = 15

# Output folders
desktop_folder = r'C:\Users\priya\OneDrive\Desktop'
project_chunks_folder = os.path.join(os.path.dirname(__file__), '..', 'seo-audit-chunks')
os.makedirs(project_chunks_folder, exist_ok=True)

# Styling
font_header = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
font_normal = Font(name="Calibri", size=9.5, color="1E293B")
font_link = Font(name="Calibri", size=9.5, color="0284C7", underline="single")
font_green = Font(name="Calibri", size=9.5, bold=True, color="15803D")

fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

col_widths = {
    1: 6, 2: 24, 3: 30, 4: 42, 5: 35, 6: 45, 7: 14, 8: 60, 9: 14,
    10: 30, 11: 45, 12: 35, 13: 42, 14: 40, 15: 65, 16: 55, 17: 60,
    18: 60, 19: 60, 20: 55, 21: 50, 22: 22, 23: 35
}

# Read header values
headers = [ws_source.cell(1, c).value for c in range(1, total_cols + 1)]

generated_files = []

for start_idx in range(1, total_data_rows + 1, chunk_size):
    end_idx = min(start_idx + chunk_size - 1, total_data_rows)
    
    file_label = f"Pages_{start_idx:02d}_to_{end_idx:02d}_SEO_Audit.xlsx"
    desktop_target = os.path.join(desktop_folder, file_label)
    project_target = os.path.join(project_chunks_folder, file_label)
    
    wb_chunk = openpyxl.Workbook()
    ws_chunk = wb_chunk.active
    ws_chunk.title = f"Pages {start_idx}-{end_idx} SEO Audit"
    ws_chunk.views.sheetView[0].showGridLines = True
    
    # Write Header
    for c_idx, h_text in enumerate(headers, 1):
        cell = ws_chunk.cell(1, c_idx, h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    
    # Write Rows
    dest_row = 2
    for src_row in range(start_idx + 1, end_idx + 2):
        fill_row = fill_zebra if dest_row % 2 == 0 else fill_white
        
        for c_idx in range(1, total_cols + 1):
            src_cell = ws_source.cell(src_row, c_idx)
            val = src_cell.value
            
            # Dynamic formula for lengths
            if c_idx == 7: # Meta Title Len
                val = f"=LEN(F{dest_row})"
            elif c_idx == 9: # Meta Desc Len
                val = f"=LEN(H{dest_row})"
                
            cell = ws_chunk.cell(dest_row, c_idx, val)
            cell.fill = fill_row
            cell.border = thin_border
            
            if c_idx in [1, 7, 9]:
                cell.alignment = align_center
                cell.font = font_bold if c_idx == 1 else font_normal
            elif c_idx in [4, 13]:
                cell.font = font_link
                cell.alignment = align_left
            elif c_idx == 22:
                cell.font = font_green
                cell.alignment = align_center
            else:
                cell.font = font_normal
                cell.alignment = align_left
                
        dest_row += 1
        
    # Column Widths
    for c_idx, w in col_widths.items():
        ws_chunk.column_dimensions[get_column_letter(c_idx)].width = w
        
    # Save
    wb_chunk.save(project_target)
    try:
        wb_chunk.save(desktop_target)
    except Exception as e:
        print(f"Could not write to desktop for {file_label}: {e}")
        
    generated_files.append((file_label, start_idx, end_idx, dest_row - 2))
    print(f"Created: {file_label} (Pages {start_idx} to {end_idx})")

print(f"\nSuccessfully generated {len(generated_files)} chunked Excel files!")
