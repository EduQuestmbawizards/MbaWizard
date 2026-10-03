import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.mbawizards.co.in"

json_path = os.path.join(os.path.dirname(__file__), 'all_blogs_data_unified.json')
all_catalog = json.loads(open(json_path, 'r', encoding='utf-8').read())

output_excel_path = os.path.join(os.path.dirname(__file__), '..', 'MBA_Wizards_All_Blogs_Complete_Catalog.xlsx')
public_excel_path = os.path.join(os.path.dirname(__file__), '..', 'public', 'MBA_Wizards_All_Blogs_Complete_Catalog.xlsx')

print(f"Loaded {len(all_catalog)} unified blog records.")

wb = openpyxl.Workbook()
wb.remove(wb.active)

# Color Scheme
NAVY_HEADER = "0F2540"
GOLD_ACCENT = "D4AF37"
ZEBRA_LIGHT = "F8FAFC"
BORDER_GRAY = "E2E8F0"
TEXT_DARK = "1E293B"

header_font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")

thin_border_side = Side(border_style="thin", color=BORDER_GRAY)
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

columns = [
    ("S.No", 8),
    ("Cluster / Category", 30),
    ("Blog Slug", 36),
    ("Blog Name / Title", 48),
    ("H1 Tag Headline", 48),
    ("Primary Keyword", 26),
    ("H2 Section Headings (Full Outline)", 65),
    ("H3 Sub-Headings", 35),
    ("Secondary Keywords / Tags", 35),
    ("Meta Title", 46),
    ("Meta Description", 55),
    ("Canonical Tag URL", 45),
    ("Lead Magnet Attached (PDF Path)", 40),
    ("Cover Image URL", 38),
    ("Author", 28),
    ("Published Date", 16),
    ("Read Time (Mins)", 16),
    ("Total Sections Count", 18),
    ("Live Blog URL", 45),
]

def write_sheet(ws, title, records):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True

    # Title Banner Row 1
    ws.merge_cells("A1:S1")
    title_cell = ws["A1"]
    title_cell.value = f"MBA WIZARDS — {title.upper()} (MASTER SEO & BLOG CATALOG)"
    title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    # Subtitle Info Row 2
    ws.merge_cells("A2:S2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Total Articles: {len(records)} | Official Canonical Domain: {BASE_URL} | Complete SEO & Content Blueprint"
    sub_cell.font = Font(name="Segoe UI", size=9.5, italic=True, color="D4AF37")
    sub_cell.fill = PatternFill(start_color="081628", end_color="081628", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Header Row 3
    for col_idx, (col_name, width) in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[3].height = 32

    # Data Rows
    for row_idx, r in enumerate(records, 4):
        s_no = row_idx - 3
        is_even = (row_idx % 2 == 0)
        row_fill = PatternFill(start_color=ZEBRA_LIGHT, end_color=ZEBRA_LIGHT, fill_type="solid") if is_even else PatternFill(fill_type=None)

        row_data = [
            s_no,
            r.get("cluster", ""),
            r.get("slug", ""),
            r.get("title", ""),
            r.get("h1", ""),
            r.get("primary_keyword", ""),
            r.get("h2_headings", ""),
            r.get("h3_headings", ""),
            r.get("secondary_keywords", ""),
            r.get("meta_title", ""),
            r.get("meta_description", ""),
            r.get("canonical_tag", ""),
            r.get("lead_magnet_pdf", ""),
            r.get("cover_image", ""),
            r.get("author", ""),
            r.get("published_at", ""),
            r.get("read_time_mins", 10),
            r.get("total_sections_count", 1),
            r.get("live_url", ""),
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=9.5, color=TEXT_DARK)
            if row_fill.fill_type:
                cell.fill = row_fill
            cell.border = cell_border

            # Alignments
            if col_idx in [1, 16, 17, 18]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in [6, 7, 8, 11]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

        # Dynamic row height for long H2 outlines
        h2_val = str(r.get("h2_headings", ""))
        num_h2_lines = h2_val.count("\n") + 1
        calculated_height = max(24, min(140, num_h2_lines * 13))
        ws.row_dimensions[row_idx].height = calculated_height

    # Auto-filter and freeze header rows
    ws.auto_filter.ref = f"A3:S{len(records) + 3}"
    ws.freeze_panes = "A4"

# 1. Master Catalog Sheet (All 198 blogs)
ws_master = wb.create_sheet()
write_sheet(ws_master, "All 198 Blogs Master Catalog", all_catalog)

# 2. MBA Admissions & Interviews (Topics 26-40)
ws_interview_26_40 = wb.create_sheet()
write_sheet(ws_interview_26_40, "MBA Interviews (26-40)", [r for r in all_catalog if r["cluster"] == "MBA Admissions & Interviews (26-40)"])

# 3. MBA Interview Strategy & Mastery (Topics 16-25)
ws_strategy_16_25 = wb.create_sheet()
write_sheet(ws_strategy_16_25, "Interview Strategy & Mastery", [r for r in all_catalog if "Interview Strategy" in r["cluster"] or "Interview Mastery" in r["cluster"]])

# 4. GMAT Score & Mock Analytics (Topics 6-15)
ws_gmat_6_15 = wb.create_sheet()
write_sheet(ws_gmat_6_15, "Score & Mock Analytics", [r for r in all_catalog if "Score Improvement" in r["cluster"] or "Mock Analytics" in r["cluster"]])

# 5. GMAT Gurgaon 50 Pillars (Topics 1-50)
ws_gurgaon_50 = wb.create_sheet()
write_sheet(ws_gurgaon_50, "GMAT Gurgaon 50 Pillars", [r for r in all_catalog if "Gurgaon 50" in r["cluster"]])

# 6. WordPress & Editorial Archive
ws_wp_archive = wb.create_sheet()
write_sheet(ws_wp_archive, "WP & Editorial Archive", [r for r in all_catalog if "Archive" in r["cluster"] or "Editorial" in r["cluster"]])

# Save Workbook
wb.save(output_excel_path)
wb.save(public_excel_path)

print(f"Successfully generated Master Excel Catalog at:\n1. {output_excel_path}\n2. {public_excel_path}")
