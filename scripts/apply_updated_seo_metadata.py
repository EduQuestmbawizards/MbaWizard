import os
import openpyxl
import re

wb = openpyxl.load_workbook('All_Website_Pages_SEO_Audit_UPDATED.xlsx', data_only=True)
ws = wb['All Website Pages SEO Audit']

page_metadata_map = {}
for r in range(2, ws.max_row + 1):
    url = ws.cell(r, 4).value
    if not url: continue
    path = url.replace("https://www.mbawizards.co.in", "").strip()
    if path == "": path = "/"
    
    title = ws.cell(r, 6).value
    desc = ws.cell(r, 8).value
    long_tails = ws.cell(r, 11).value
    
    page_metadata_map[path] = {
        "title": title,
        "desc": desc,
        "long_tails": long_tails,
        "sno": ws.cell(r, 1).value,
        "name": ws.cell(r, 3).value
    }

print(f"Loaded {len(page_metadata_map)} page mappings from updated Excel.")

# Let's check matching files in src/app
app_dir = 'src/app'
updated_count = 0

for path, data in page_metadata_map.items():
    if path.startswith("/blogs/"):
        # These are blogs, handled in gmat-gurgaon-blogs.ts
        continue
    
    if path == "/":
        file_path = os.path.join(app_dir, '(home)', 'page.tsx')
    else:
        folder = path.lstrip('/')
        file_path = os.path.join(app_dir, folder, 'page.tsx')
        if not os.path.exists(file_path):
            file_path = os.path.join(app_dir, folder, 'page.ts')
    
    if os.path.exists(file_path):
        content = open(file_path, 'r', encoding='utf-8').read()
        
        # Check if metadata is exported
        new_title = data["title"].replace('"', '\\"')
        new_desc = data["desc"].replace('"', '\\"')
        
        # Update metadata object if present
        updated = False
        if "export const metadata" in content:
            # Replace title
            content = re.sub(r'title:\s*["\'`][^"\'`]+["\'`]', f'title: "{new_title}"', content, count=1)
            # Replace description
            content = re.sub(r'description:\s*["\'`][^"\'`]+["\'`]', f'description: "{new_desc}"', content, count=1)
            updated = True
        
        if updated:
            open(file_path, 'w', encoding='utf-8').write(content)
            updated_count += 1
            print(f"Updated {file_path} -> {data['title'][:40]}...")
        else:
            print(f"No metadata export in {file_path}")
    else:
        print(f"File not found for path: {path} -> {file_path}")

print(f"\nTotal page files updated: {updated_count}")
